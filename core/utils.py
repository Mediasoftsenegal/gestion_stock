"""Fonctions utilitaires transverses."""

import datetime as dt
import functools
import unicodedata
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

from django.db import IntegrityError, OperationalError, transaction
from django.utils import timezone


# --------------------------------------------------------------------------- #
#  Parsing tolerant de valeurs issues de fichiers (Excel/CSV)
# --------------------------------------------------------------------------- #
def normaliser_texte(valeur):
    """Normalise un texte pour comparaison tolerante (accents/casse/espaces)."""
    valeur = (valeur or "").strip()
    valeur = unicodedata.normalize("NFD", valeur).encode("ascii", "ignore").decode("ascii")
    valeur = valeur.replace("_", " ").replace("-", " ").replace(".", " ")
    return " ".join(valeur.upper().split())


def parser_decimal(valeur, defaut=Decimal("0")):
    valeur = str(valeur or "").strip()
    if not valeur:
        return defaut
    valeur = valeur.replace("\xa0", "").replace(" ", "").replace(",", ".")
    valeur = valeur.replace("FCFA", "").replace("F", "") if valeur.upper().endswith("FCFA") else valeur
    try:
        return Decimal(valeur)
    except (InvalidOperation, ValueError):
        return defaut


def parser_entier(valeur, defaut=0):
    dec = parser_decimal(valeur, defaut=None)
    if dec is None:
        return defaut
    try:
        return int(dec.to_integral_value(rounding=ROUND_HALF_UP))
    except (InvalidOperation, ValueError):
        return defaut


def parser_booleen(valeur, defaut=True):
    texte = normaliser_texte(valeur)
    if not texte:
        return defaut
    if texte in {"1", "OUI", "O", "TRUE", "VRAI", "ACTIF", "YES", "Y", "X"}:
        return True
    if texte in {"0", "NON", "N", "FALSE", "FAUX", "INACTIF", "NO"}:
        return False
    return defaut


def reessayer_si_conflit_bd(tentatives=3):
    """Decorateur de vue : rejoue l'operation entiere (nouvelle transaction)
    si la base leve un deadlock transitoire (`OperationalError`), plutot que
    de renvoyer une erreur 500 pour un conflit sans consequence.

    A utiliser sur les vues qui creent des enregistrements references
    (commande, facture, approvisionnement, inventaire...), sous la
    decoration `login_required`/`permission_required`/`require_POST` pour
    ne rejouer que la logique metier, pas les verifications d'acces.
    """

    def decorateur(vue):
        @functools.wraps(vue)
        def wrapper(*args, **kwargs):
            for essai in range(1, tentatives + 1):
                try:
                    return vue(*args, **kwargs)
                except OperationalError:
                    if essai == tentatives:
                        raise

        return wrapper

    return decorateur


def generer_reference(prefixe, model, champ="reference"):
    """Genere une reference du type PREFIXE-ANNEE-0001, garantie sans doublon
    meme en cas de creations concurrentes (plusieurs utilisateurs sur le
    reseau qui valident au meme moment) : incremente un compteur verrouille
    (`SELECT ... FOR UPDATE`) plutot que de relire/recalculer la derniere
    reference existante, ce qui serait sujet a une condition de course.

    Sous tres forte concurrence (creations simultanees d'un tout nouveau
    compteur, ex. premiere commande de l'annee), MySQL peut lever un
    deadlock (`OperationalError`) : c'est une erreur propre et sans danger
    a rejouer (rien n'a ete valide) ; voir `core.utils.reessayer_si_conflit_bd`
    pour l'encadrer au niveau vue/formulaire.
    """
    from core.models import SequenceCompteur

    annee = timezone.now().year
    cle = f"{prefixe}-{annee}"
    base = f"{cle}-"

    with transaction.atomic():
        try:
            compteur = SequenceCompteur.objects.select_for_update().get(cle=cle)
        except SequenceCompteur.DoesNotExist:
            # Premiere utilisation de ce compteur : on part du plus grand
            # numero deja utilise (donnees existantes) pour ne pas regenerer
            # une reference deja prise.
            numero_initial = 0
            dernier = (
                model.objects.filter(**{f"{champ}__startswith": base})
                .order_by(f"-{champ}")
                .values_list(champ, flat=True)
                .first()
            )
            if dernier:
                try:
                    numero_initial = int(dernier.rsplit("-", 1)[1])
                except (ValueError, IndexError):
                    numero_initial = model.objects.filter(
                        **{f"{champ}__startswith": base}
                    ).count()
            try:
                with transaction.atomic():
                    compteur = SequenceCompteur.objects.create(
                        cle=cle, dernier_numero=numero_initial
                    )
            except IntegrityError:
                # Une transaction concurrente vient de le creer : on la
                # relit, verrouillee, plutot que d'echouer.
                compteur = SequenceCompteur.objects.select_for_update().get(cle=cle)

        compteur.dernier_numero += 1
        compteur.save(update_fields=["dernier_numero"])
        numero = compteur.dernier_numero

    return f"{base}{numero:04d}"


def generer_code(prefixe, model, champ="code", chiffres=5):
    """Genere un code sequentiel permanent (sans annee), du type PREFIXE00001,
    garanti sans doublon meme en cas de creations concurrentes (meme
    verrouillage que `generer_reference`). A la difference de celle-ci, la
    numerotation n'est jamais remise a zero : il s'agit d'un catalogue
    permanent (articles...), pas de documents annuels.
    """
    from core.models import SequenceCompteur

    cle = f"CODE-{prefixe}"

    with transaction.atomic():
        try:
            compteur = SequenceCompteur.objects.select_for_update().get(cle=cle)
        except SequenceCompteur.DoesNotExist:
            numero_initial = 0
            dernier = (
                model.objects.filter(**{f"{champ}__startswith": prefixe})
                .order_by(f"-{champ}")
                .values_list(champ, flat=True)
                .first()
            )
            if dernier:
                try:
                    numero_initial = int(dernier[len(prefixe):])
                except ValueError:
                    numero_initial = model.objects.filter(
                        **{f"{champ}__startswith": prefixe}
                    ).count()
            try:
                with transaction.atomic():
                    compteur = SequenceCompteur.objects.create(
                        cle=cle, dernier_numero=numero_initial
                    )
            except IntegrityError:
                compteur = SequenceCompteur.objects.select_for_update().get(cle=cle)

        compteur.dernier_numero += 1
        compteur.save(update_fields=["dernier_numero"])
        numero = compteur.dernier_numero

    return f"{prefixe}{numero:0{chiffres}d}"


PERIODES = [
    ("JOUR", "Jour"),
    ("SEMAINE", "Semaine"),
    ("QUINZAINE", "Quinzaine"),
    ("MOIS", "Mois"),
    ("TRIMESTRE", "Trimestre"),
    ("SEMESTRE", "Semestre"),
    ("ANNEE", "Annee"),
]


def bornes_periode(periode, reference=None):
    """Retourne (debut, fin) inclusifs pour une periode donnee.

    `reference` est une date situee dans la periode (defaut : aujourd'hui).
    """
    ref = reference or timezone.localdate()

    if periode == "JOUR":
        return ref, ref
    if periode == "SEMAINE":
        debut = ref - dt.timedelta(days=ref.weekday())
        return debut, debut + dt.timedelta(days=6)
    if periode == "QUINZAINE":
        debut = ref.replace(day=1) if ref.day <= 15 else ref.replace(day=16)
        if ref.day <= 15:
            return debut, ref.replace(day=15)
        fin = _dernier_jour_mois(ref)
        return debut, fin
    if periode == "MOIS":
        return ref.replace(day=1), _dernier_jour_mois(ref)
    if periode == "TRIMESTRE":
        premier_mois = 3 * ((ref.month - 1) // 3) + 1
        debut = ref.replace(month=premier_mois, day=1)
        fin_mois = premier_mois + 2
        return debut, _dernier_jour_mois(ref.replace(month=fin_mois, day=1))
    if periode == "SEMESTRE":
        premier_mois = 1 if ref.month <= 6 else 7
        debut = ref.replace(month=premier_mois, day=1)
        fin = ref.replace(month=premier_mois + 5, day=1)
        return debut, _dernier_jour_mois(fin)
    if periode == "ANNEE":
        return ref.replace(month=1, day=1), ref.replace(month=12, day=31)

    raise ValueError(f"Periode inconnue : {periode}")


def _dernier_jour_mois(date):
    if date.month == 12:
        suivant = date.replace(year=date.year + 1, month=1, day=1)
    else:
        suivant = date.replace(month=date.month + 1, day=1)
    return suivant - dt.timedelta(days=1)


# --------------------------------------------------------------------------- #
#  Conversion d'un montant en toutes lettres (francais)
# --------------------------------------------------------------------------- #
_UNITES = [
    "zéro", "un", "deux", "trois", "quatre", "cinq", "six", "sept", "huit", "neuf",
    "dix", "onze", "douze", "treize", "quatorze", "quinze", "seize",
    "dix-sept", "dix-huit", "dix-neuf",
]
_DIZAINES = {2: "vingt", 3: "trente", 4: "quarante", 5: "cinquante", 6: "soixante", 8: "quatre-vingt"}


def _deux_chiffres_en_lettres(n):
    """n entre 0 et 99."""
    if n < 20:
        return _UNITES[n]
    if n < 70:
        d, u = divmod(n, 10)
        mot = _DIZAINES[d]
        if u == 0:
            return mot
        if u == 1:
            return f"{mot} et un"
        return f"{mot}-{_UNITES[u]}"
    if n < 80:
        reste = n - 60
        if reste == 11:
            return "soixante et onze"
        return f"soixante-{_UNITES[reste]}"
    reste = n - 80
    if reste == 0:
        return "quatre-vingts"
    return f"quatre-vingt-{_UNITES[reste]}"


def _trois_chiffres_en_lettres(n):
    """n entre 0 et 999."""
    centaines, reste = divmod(n, 100)
    parts = []
    if centaines > 0:
        if centaines == 1:
            parts.append("cent")
        else:
            parts.append(f"{_UNITES[centaines]} cent" + ("s" if reste == 0 else ""))
    if reste > 0:
        parts.append(_deux_chiffres_en_lettres(reste))
    return " ".join(parts) if parts else "zéro"


def nombre_en_lettres(n):
    """Convertit un entier (positif, negatif ou nul) en toutes lettres (francais)."""
    n = int(n)
    if n == 0:
        return "zéro"
    if n < 0:
        return "moins " + nombre_en_lettres(-n)

    tranches = []
    reste = n
    while reste > 0:
        tranches.append(reste % 1000)
        reste //= 1000

    mots = []
    for i in range(len(tranches) - 1, -1, -1):
        valeur = tranches[i]
        if valeur == 0:
            continue
        if i == 1:
            mots.append("mille" if valeur == 1 else f"{_trois_chiffres_en_lettres(valeur)} mille")
        elif i >= 2:
            echelle = "million" if i == 2 else "milliard"
            mots.append(f"{_trois_chiffres_en_lettres(valeur)} {echelle}{'s' if valeur > 1 else ''}")
        else:
            mots.append(_trois_chiffres_en_lettres(valeur))
    return " ".join(mots)


def montant_en_lettres(montant, devise_singulier="franc CFA", devise_pluriel="francs CFA"):
    """Convertit un montant en toutes lettres, ex : 'Cent vingt-cinq mille francs CFA'."""
    entier = int(Decimal(montant).to_integral_value(rounding=ROUND_HALF_UP))
    devise = devise_singulier if abs(entier) <= 1 else devise_pluriel
    mots = nombre_en_lettres(entier)
    return f"{mots[0].upper()}{mots[1:]} {devise}"


# --------------------------------------------------------------------------- #
#  Import / export Excel (.xlsx, via openpyxl)
# --------------------------------------------------------------------------- #
def generer_modele_xlsx(colonnes, exemples=None, feuille="Import"):
    """Construit un classeur Excel vierge (en-tetes + lignes d'exemple
    facultatives, grisees) pret a etre rempli puis reimporte.

    `colonnes` : liste de libelles d'en-tete, dans l'ordre attendu.
    `exemples` : liste de tuples (une ligne d'exemple par tuple), facultatif.
    Retourne les octets du fichier .xlsx.
    """
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    classeur = Workbook()
    feuille_active = classeur.active
    feuille_active.title = feuille

    feuille_active.append(list(colonnes))
    for cellule in feuille_active[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="9C2B2B")

    for ligne in exemples or []:
        feuille_active.append(list(ligne))
    for row in feuille_active.iter_rows(min_row=2):
        for cellule in row:
            cellule.font = Font(italic=True, color="808080")

    for index, titre in enumerate(colonnes, start=1):
        feuille_active.column_dimensions[feuille_active.cell(row=1, column=index).column_letter].width = max(
            14, len(titre) + 2
        )

    tampon = BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


def lire_xlsx(fichier):
    """Lit un classeur Excel (.xlsx/.xlsm) et retourne (entetes, lignes).

    `lignes` est une liste de dictionnaires {en-tete (str, en majuscules,
    sans accents) : valeur brute (str)}. La premiere ligne non vide du
    classeur est consideree comme l'en-tete.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depend de l'environnement
        raise ValueError(
            "La lecture des fichiers Excel necessite openpyxl (pip install openpyxl)."
        ) from exc

    def normaliser_entete(valeur):
        valeur = unicodedata.normalize("NFD", str(valeur or "")).encode("ascii", "ignore").decode("ascii")
        return " ".join(valeur.upper().split())

    classeur = load_workbook(fichier, data_only=True, read_only=True)
    feuille = classeur.active
    iterateur = feuille.iter_rows(values_only=True)

    entetes = []
    for brut in iterateur:
        if brut and any(c is not None and str(c).strip() for c in brut):
            entetes = [normaliser_entete(c) for c in brut]
            break

    lignes = []
    for brut in iterateur:
        if not brut or not any(c is not None and str(c).strip() for c in brut):
            continue
        valeurs = ["" if c is None else str(c).strip() for c in brut]
        valeurs += [""] * (len(entetes) - len(valeurs))
        lignes.append(dict(zip(entetes, valeurs)))

    return entetes, lignes
